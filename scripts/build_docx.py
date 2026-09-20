"""Build DOCX copies of workbooks and the report for the submission zip."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent


def add_md_ish(doc: Document, text: str) -> None:
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("|") and "---" not in line:
            doc.add_paragraph(line, style="Normal")
        else:
            p = doc.add_paragraph(line)
            for run in p.runs:
                run.font.size = Pt(11)


def convert_markdown(src: Path, dest: Path) -> None:
    doc = Document()
    add_md_ish(doc, src.read_text(encoding="utf-8"))
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def report_with_appendix() -> None:
    src = ROOT / "docs" / "MaharshiGogula_Capstone_Report.md"
    dest = ROOT / "docs" / "MaharshiGogula_Capstone_Report.docx"
    doc = Document()
    add_md_ish(doc, src.read_text(encoding="utf-8"))
    metrics_path = ROOT / "evaluation" / "results" / "metrics.json"
    if metrics_path.exists():
        doc.add_heading("Appendix A — metrics.json (validation run)", 1)
        metrics = json.loads(metrics_path.read_text(encoding="utf-8"))
        doc.add_paragraph(json.dumps(metrics, indent=2)[:12000])
    doc.save(dest)


def main() -> None:
    wb = ROOT / "workbooks"
    for md in sorted(wb.glob("*.md")):
        convert_markdown(md, wb / (md.stem + ".docx"))
    report_with_appendix()
    print("wrote docx files")


if __name__ == "__main__":
    main()
